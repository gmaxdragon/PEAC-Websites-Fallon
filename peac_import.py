"""Bounded spreadsheet preview and transactional import. No uploads are saved.
CSV/TSV are built-in; optional XLSX reading uses openpyxl + defusedxml.
Aggregate imports remain non-identifying. A separate authenticated named-weekly dataset is supported for staff-only roster logging. Preview is not a save.
"""
from __future__ import annotations
import base64
import binascii
import csv
import hashlib
import io
import json
import re
import zipfile
from datetime import date, datetime
from peac_core import APIError, MAX_COUNT, count, normalize_campaign, object_payload, recorded_date, text, validate_count_entry
from peac_people import current_people_rows, insert_person_entry, validate_person_record

MAX_ROWS = 5000
MAX_TEXT = 1_500_000
MAX_BINARY = 2_000_000
MAX_HTTP = 3_000_000
ALIASES = {
    'date':'date','day':'date','week':'date','week start':'date',
    'count':'count','verified':'count','verified compliments':'count','compliments':'count',
    'trash':'trash_count','trash count':'trash_count','trash compliments':'trash_count',
    'grade':'grade','grade level':'grade','note':'note','notes':'note',
    'name':'name','student':'name','student name':'name','person':'name','campaign':'name','campaign name':'name','type':'type','platform':'platform',
    'description':'description','impressions':'impressions','engagements':'engagements',
    'compliments attributed':'compliments_attributed','attributed compliments':'compliments_attributed',
}
COUNT_FIELDS = {'date','count','trash_count','grade','note'}
PEOPLE_FIELDS = {'date','name','count','trash_count','grade','note'}
CAMPAIGN_FIELDS = {'date','name','type','platform','description','impressions','engagements','compliments_attributed','note'}


def fingerprint(dataset, records):
    rows = sorted(records, key=lambda r: json.dumps(r,sort_keys=True,ensure_ascii=False))
    return hashlib.sha256(json.dumps([dataset,rows],sort_keys=True,ensure_ascii=False,allow_nan=False).encode()).hexdigest()


def parse_number(value, field, optional=False):
    if value in (None,''):
        if optional: return 0
        raise APIError(f'{field} is blank. Missing is not zero.')
    if type(value) is int: return count(value,field)
    if type(value) is float and value.is_integer(): return count(int(value),field)
    if not isinstance(value,str): raise APIError(f'{field} must be a whole number.')
    value=value.strip()
    if not re.fullmatch(r'(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.0+)?',value):
        raise APIError(f'{field} must be a non-negative whole number, not a formula.')
    return count(int(value.replace(',','').split('.')[0]),field)


def parse_date(value, today, date_format):
    if isinstance(value,datetime): value=value.date()
    if isinstance(value,date): return recorded_date(value.isoformat(),today)
    if not isinstance(value,str): raise APIError('Date is required. Use a formatted Excel date or YYYY-MM-DD.')
    value=value.strip()
    if re.fullmatch(r'\d{4}-\d{2}-\d{2}',value): return recorded_date(value,today)
    formats={'mdy':'%m/%d/%Y','dmy':'%d/%m/%Y'}
    if date_format in formats:
        try: return recorded_date(datetime.strptime(value,formats[date_format]).date().isoformat(),today)
        except ValueError: pass
    raise APIError('Date must be YYYY-MM-DD, or match the explicitly selected slash-date format.')


def read_matrix(payload, sheets_only=False):
    fmt=payload.get('format','csv')
    if fmt not in {'csv','tsv','xlsx'}: raise APIError('Choose CSV, pasted cells, or XLSX. XLS/XLSM are not supported.')
    sheet_names=[]
    if fmt in {'csv','tsv'}:
        raw=payload.get('text')
        if not isinstance(raw,str) or not raw.strip(): raise APIError('Choose a file or paste cells, including the header row.')
        if len(raw.encode('utf-8')) > MAX_TEXT: raise APIError('Spreadsheet text is too large (1.5 MB maximum).',413)
        raw=raw.lstrip('\ufeff')
        delimiter='\t' if fmt=='tsv' else ','
        try:
            # Google Sheets/Excel clipboard uses tabs. CSV can use commas or semicolons.
            if fmt=='csv':
                try: delimiter=csv.Sniffer().sniff(raw[:16000],delimiters=',;\t').delimiter
                except csv.Error: pass
            matrix=[]
            reader=csv.reader(io.StringIO(raw,newline=''),delimiter=delimiter,strict=True)
            for row in reader:
                if len(row)>16: raise APIError('At most 16 columns. Use a sanitized aggregate-only sheet.')
                if any(str(v).strip() for v in row): matrix.append(row)
                if len(matrix)>MAX_ROWS+1: raise APIError('At most 5,000 records per import.',413)
        except csv.Error as exc: raise APIError('Invalid CSV quoting or oversized cell.') from exc
    else:
        encoded=payload.get('content_base64')
        if not isinstance(encoded,str) or len(encoded)>2_700_000: raise APIError('XLSX file is too large (2 MB maximum).',413)
        try: raw=base64.b64decode(encoded,validate=True)
        except (ValueError,binascii.Error): raise APIError('Invalid XLSX upload.')
        if len(raw)>MAX_BINARY: raise APIError('XLSX file exceeds 2 MB.',413)
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as archive:
                files=archive.infolist()
                if len(files)>300 or sum(f.file_size for f in files)>20_000_000:
                    raise APIError('Workbook is too complex or expands beyond 20 MB.',413)
                if any(f.flag_bits & 1 for f in files): raise APIError('Encrypted workbooks are not supported.')
                if any('vbaproject' in f.filename.lower() or 'externallinks/' in f.filename.lower() for f in files):
                    raise APIError('Use a clean .xlsx file without macros or external workbook links.')
        except zipfile.BadZipFile: raise APIError('File is not a valid XLSX workbook.')
        try:
            import defusedxml  # Required, not optional, for untrusted workbook XML.
            from openpyxl import load_workbook
        except ImportError: raise APIError('XLSX support is not installed. Use CSV/paste, or install requirements-excel.txt.',415)
        try:
            book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
            try:
                sheet_names=book.sheetnames
                if sheets_only: return [],sheet_names
                if not sheet_names: raise APIError('Workbook contains no worksheets.')
                requested=payload.get('sheet') or sheet_names[0]
                if requested not in sheet_names: raise APIError('Selected worksheet does not exist.')
                sheet=book[requested]
                if (sheet.max_column or 0)>16 or (sheet.max_row or 0)>MAX_ROWS+1:
                    raise APIError('Selected sheet exceeds 16 columns or 5,000 records. Copy only aggregate data into a new workbook.')
                matrix=[]
                for row in sheet.iter_rows():
                    if any(cell.data_type=='f' for cell in row):
                        raise APIError('Formula cells are not imported. Copy and paste values only into a clean sheet first.')
                    values=[cell.value for cell in row]
                    while values and values[-1] is None: values.pop()
                    if values: matrix.append(values)
                    if len(matrix)>MAX_ROWS+1: raise APIError('At most 5,000 records per import.',413)
            finally: book.close()
        except APIError: raise
        except Exception as exc: raise APIError('Workbook could not be read. Save a clean XLSX or use CSV/paste.') from exc
    if len(matrix)<2: raise APIError('Include a header row and at least one data row.')
    return matrix,sheet_names


def parse_upload(payload,today):
    object_payload(payload)
    allowed={'format','text','content_base64','sheet','dataset','date_format','mode','source'}
    if set(payload)-allowed: raise APIError('Unknown spreadsheet upload fields.')
    dataset=payload.get('dataset','counts')
    if dataset not in {'counts','people','campaigns'}: raise APIError('Choose aggregate Counts, Named weekly log, or Campaigns.')
    mode=payload.get('mode','new_only')
    if mode not in {'new_only','skip_existing'}: raise APIError('Choose new rows only or skip existing records.')
    date_format=payload.get('date_format','iso')
    if date_format not in {'iso','mdy','dmy'}: raise APIError('Unknown date format.')
    source=text(payload.get('source','Spreadsheet import'),'source',120) or 'Spreadsheet import'
    matrix,sheets=read_matrix(payload)
    fields=[]
    for header in matrix[0]:
        normalized=re.sub(r'[_\-\s]+',' ',str(header or '').strip().lower())
        field=ALIASES.get(normalized)
        allowed_fields = COUNT_FIELDS if dataset=='counts' else PEOPLE_FIELDS if dataset=='people' else CAMPAIGN_FIELDS
        if not field or field not in allowed_fields:
            if dataset=='people':
                raise APIError('Unsupported column header. Named weekly logs accept only week_start/date, name, grade, count, trash_count, and note.')
            raise APIError('Unsupported column header. Use only the template fields; aggregate analytics do not accept student names.')
        if field in fields: raise APIError('Two columns map to the same field. Keep only one.')
        fields.append(field)
    required={'date','count'} if dataset=='counts' else {'date','name','count','grade'} if dataset=='people' else {'date','name'}
    if not required <= set(fields): raise APIError('Required columns: '+', '.join(sorted(required)))
    records=[]; errors=[]; record_rows=[]
    for index,values in enumerate(matrix[1:],start=2):
        try:
            if len(values)>len(fields): raise APIError('Too many cells for the header row.')
            value=dict(zip(fields,values+['']*(len(fields)-len(values))))
            value['date']=parse_date(value.get('date'),today,date_format)
            if dataset in {'counts','people'}:
                value['count']=parse_number(value.get('count'),'verified count' if dataset=='counts' else 'compliments')
                value['trash_count']=parse_number(value.get('trash_count'),'trash count',True)
                grade=str(value.get('grade') or 'unassigned').strip().lower().replace('th grade','').replace('grade ','')
                value['grade']=grade if grade!='school-wide' else 'unassigned'
                value['note']=value.get('note') or ''
                if dataset=='counts':
                    record=validate_count_entry(value,today)
                else:
                    value['week_start']=value.pop('date')
                    record=validate_person_record(value,today)
            else:
                for field in ('impressions','engagements','compliments_attributed'): value[field]=parse_number(value.get(field),field,True)
                if 'note' in value: value['notes']=value.pop('note') or ''
                for field in ('type','platform'): value[field]=value.get(field) or 'Other'
                if 'description' in value: value['description']=value['description'] or ''
                normalized=normalize_campaign(value,today)
                record={k:v for k,v in normalized.items() if k not in {'id','created_at','updated_at'}}
            records.append(record); record_rows.append(index)
        except APIError as exc:
            errors.append({'row':index,'message':str(exc)})
    return {'dataset':dataset,'mode':mode,'source':source,'records':records,'errors':errors[:100],
            'error_count':len(errors),'row_count':len(matrix)-1,'sheets':sheets,'record_rows':record_rows,
            'sheet':payload.get('sheet') or (sheets[0] if sheets else '')}


def overlaps(record,existing,dataset):
    if dataset=='counts':
        return any(r['date']==record['date'] and (r['grade']==record['grade'] or 'unassigned' in {r['grade'],record['grade']}) for r in existing)
    if dataset=='people':
        return any(r['date']==record['date'] and r['grade']==record['grade'] and r['name'].strip().casefold()==record['name'].strip().casefold() for r in existing)
    return any(r['date']==record['date'] and r['name'].casefold()==record['name'].casefold() for r in existing)


def import_plan(records,dataset,mode,snapshot):
    existing=snapshot['entries'] if dataset=='counts' else snapshot.get('people',[]) if dataset=='people' else snapshot['campaigns']
    previous=[]; included=[]; skipped=[]; duplicates=[]
    for index,r in enumerate(records,start=2):
        if overlaps(r,previous,dataset): duplicates.append(index)
        previous.append(r)
        if overlaps(r,existing,dataset): skipped.append(index)
        else: included.append(r)
    digest=fingerprint(dataset,records)
    already=any(b['fingerprint']==digest for b in snapshot.get('imports',[]))
    return {'included':included,'existing_rows':skipped,'duplicate_rows':duplicates,'fingerprint':digest,
            'already_imported':already,
            'can_import':bool(included) and not duplicates and not already and (mode=='skip_existing' or not skipped)}


def preview(service,payload):
    result=parse_upload(payload,service.today_fn())
    snapshot=service.store.snapshot()
    if result['dataset']=='people': snapshot['people']=current_people_rows(service)
    plan=import_plan(result['records'],result['dataset'],result['mode'],snapshot)
    result.update({k:v for k,v in plan.items() if k!='included'})
    result['existing_rows']=[result['record_rows'][i-2] for i in plan['existing_rows']]
    result['duplicate_rows']=[result['record_rows'][i-2] for i in plan['duplicate_rows']]
    result['can_import']=plan['can_import'] and not result['error_count']
    result['revision']=snapshot['revision']
    result['new_rows']=len(plan['included'])
    result['totals']={k:sum(r.get(k,0) for r in plan['included']) for k in ('count','trash_count','impressions','engagements','compliments_attributed')}
    return result


def commit(service,payload,key):
    object_payload(payload)
    if set(payload)-{'dataset','records','mode','source','revision','fingerprint','confirm'}: raise APIError('Unknown import fields.')
    dataset=payload.get('dataset'); mode=payload.get('mode')
    if dataset not in {'counts','people','campaigns'} or mode not in {'new_only','skip_existing'}: raise APIError('Invalid import options.')
    if payload.get('confirm')!='IMPORT PREVIEWED ROWS': raise APIError('Preview and confirm the import first.')
    if type(payload.get('revision')) is not int: raise APIError('A current preview revision is required.')
    records=payload.get('records')
    if not isinstance(records,list) or not 1<=len(records)<=MAX_ROWS: raise APIError('Import 1 to 5,000 records.')
    normalized=[]
    for row in records:
        if dataset=='counts': normalized.append(validate_count_entry(row,service.today_fn()))
        elif dataset=='people': normalized.append(validate_person_record(row,service.today_fn()))
        else:
            c=normalize_campaign(row,service.today_fn());normalized.append({k:v for k,v in c.items() if k not in {'id','created_at','updated_at'}})
    digest=fingerprint(dataset,normalized)
    if digest!=payload.get('fingerprint'): raise APIError('Preview contents changed. Preview again.',409)
    source=text(payload.get('source','Spreadsheet import'),'source',120) or 'Spreadsheet import'
    def operation(db):
        from peac_core import now_iso
        import uuid
        revision=int(db.execute("SELECT value FROM metadata WHERE key='revision'").fetchone()[0])
        if revision!=payload['revision']: raise APIError('The database changed since preview. Preview again before importing.',409)
        snapshot={'entries':[dict(r) for r in db.execute('SELECT * FROM entries')],
                  'campaigns':[json.loads(r[0]) for r in db.execute('SELECT payload FROM campaigns')],
                  'imports':[dict(r) for r in db.execute('SELECT fingerprint FROM import_batches')]}
        if dataset=='people':
            snapshot['people']=[dict(r) for r in db.execute(
                """SELECT p.entry_id AS id,p.person_name AS name,e.date,e.grade,e.count,e.trash_count
                     FROM compliment_people p JOIN entries e ON e.id=p.entry_id"""
            )]
        plan=import_plan(normalized,dataset,mode,snapshot)
        if plan['already_imported']: raise APIError('This exact batch was already imported. Nothing changed.',409)
        if plan['duplicate_rows']: raise APIError('Duplicate date/grade or campaign keys in the batch. Nothing changed.',409)
        if not plan['can_import']: raise APIError('No new rows, or existing records conflict. Nothing changed.',409)
        identities=[]
        for value in plan['included']:
            if dataset=='counts': identities.append(service.store.insert_entry(db,value))
            elif dataset=='people': identities.append(insert_person_entry(service,db,value,'spreadsheet-import'))
            else:
                value=normalize_campaign(value,service.today_fn());identities.append(value['id'])
                db.execute('INSERT INTO campaigns VALUES(?,?)',(value['id'],json.dumps(value)))
        identity=str(uuid.uuid4())
        db.execute('INSERT INTO import_batches VALUES(?,?,?,?,?,?,?,?)',
                   (identity,digest,dataset,source,len(identities),len(plan['existing_rows']),json.dumps(identities),now_iso()))
        return {'batch_id':identity,'imported':len(identities),'skipped':len(plan['existing_rows']),'dataset':dataset}
    return service.store.mutate(key,'spreadsheet-import',payload,operation)


def list_sheets(payload):
    object_payload(payload)
    if payload.get("format")!="xlsx": raise APIError("Worksheet selection applies to XLSX only.")
    _,names=read_matrix(payload,sheets_only=True)
    return {"sheets":names}
